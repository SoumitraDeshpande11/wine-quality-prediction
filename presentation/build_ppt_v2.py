from pathlib import Path
import json
import math
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'presentation' / 'Wine_Quality_Prediction_BDA.pptx'
DATA = pd.read_csv(ROOT / 'data' / 'winequality-red.csv', sep=';')
MET = json.loads((ROOT / 'artifacts' / 'metrics.json').read_text())['models']
FI = pd.read_csv(ROOT / 'artifacts' / 'feature_importance.csv')
W, H = 13.333, 7.5
C = dict(bg='17171A', bg2='202025', panel='29282D', panel2='343039', wine='822D43', wine2='A63C55', blush='CD8D99', gold='E8C68A', cream='F7F1E7', muted='B8B0A7', quiet='817C79', line='4A4448', white='FFFFFF', darkwine='551D30')
FONT='Arial'
SERIF='Georgia'

def rgb(v):
    v=C.get(v,v).lstrip('#')
    return RGBColor(*(int(v[i:i+2],16) for i in (0,2,4)))

def rect(s,x,y,w,h,fill='panel',stroke=None,radius=False):
    z=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    z.fill.solid(); z.fill.fore_color.rgb=rgb(fill)
    if stroke: z.line.color.rgb=rgb(stroke); z.line.width=Pt(.75)
    else: z.line.fill.background()
    return z

def line(s,x1,y1,x2,y2,color='line',width=1):
    z=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    z.line.color.rgb=rgb(color); z.line.width=Pt(width); return z

def oval(s,x,y,w,h,fill='wine',stroke=None):
    z=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(w),Inches(h));z.fill.solid();z.fill.fore_color.rgb=rgb(fill)
    if stroke:z.line.color.rgb=rgb(stroke);z.line.width=Pt(1)
    else:z.line.fill.background()
    return z

def txt(s,x,y,w,h,t,size=16,color='cream',bold=False,font=FONT,align=PP_ALIGN.LEFT,ital=False,anchor=MSO_ANCHOR.TOP):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));f=sh.text_frame;f.clear();f.word_wrap=True
    f.margin_left=f.margin_right=f.margin_top=f.margin_bottom=0;f.vertical_anchor=anchor
    p=f.paragraphs[0];p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(0)
    r=p.add_run();r.text=str(t);r.font.name=font;r.font.size=Pt(size);r.font.bold=bold;r.font.italic=ital;r.font.color.rgb=rgb(color)
    return sh

def link(s,x,y,w,h,t,url,size=12,color='gold'):
    sh=txt(s,x,y,w,h,t,size,color,True)
    sh.text_frame.paragraphs[0].runs[0].hyperlink.address=url
    return sh

def base(prs,num,kicker=None):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb('bg')
    rect(s,0,0,.12,H,'wine')
    line(s,.55,7.09,12.8,7.09,'line',.8)
    txt(s,.58,7.17,6,.18,'TEAM BDA  /  WINE QUALITY PREDICTION',8.5,'quiet',True)
    txt(s,12.24,7.15,.5,.2,f'{num:02d}',9,'gold',True,align=PP_ALIGN.RIGHT)
    if kicker: txt(s,.58,.37,7,.22,kicker.upper(),10,'gold',True)
    return s

def heading(s,title,sub=None):
    txt(s,.58,.75,12.15,.62,title,30,'cream',False,SERIF)
    if sub: txt(s,.61,1.43,11.9,.45,sub,13,'muted')

def bar(s,x,y,w,h,value,maxval,fill='wine2',bg='panel'):
    rect(s,x,y,w,h,bg);rect(s,x,y,w*max(0,value)/maxval,h,fill)

def node(s,cx,cy,label,fill='gold',textcolor='bg'):
    oval(s,cx-.16,cy-.16,.32,.32,fill)
    txt(s,cx-.13,cy-.1,.26,.18,label,8.5,textcolor,True,align=PP_ALIGN.CENTER)

def make():
    prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H)

    # 01 — Title: editorial, intentionally asymmetric.
    s=base(prs,1)
    rect(s,8.55,0,4.78,7.09,'wine')
    for i in range(8): line(s,8.73+i*.43,.42,12.85,5.3+i*.18,'darkwine',.7)
    oval(s,9.45,1.16,2.42,2.42,'darkwine','gold')
    oval(s,10.04,1.74,1.25,1.25,'wine','gold')
    line(s,10.67,2.96,10.67,4.28,'gold',3)
    line(s,9.92,4.29,11.42,4.29,'gold',3)
    txt(s,.62,.64,5.9,.25,'MACHINE LEARNING FUNDAMENTALS  ·  GROUP 14',11,'gold',True)
    txt(s,.58,1.36,7.65,2.55,'WINE\nQUALITY',49,'cream',False,SERIF)
    txt(s,.64,4.18,6.5,.52,'PREDICTION',25,'gold',True)
    line(s,.64,4.9,6.94,4.9,'gold',1.5)
    txt(s,.65,5.18,6.45,.62,'From red-wine chemistry to an interpretable quality estimate',16,'muted')
    txt(s,9.08,5.28,3.54,.3,'TEAM BDA',12,'gold',True)
    txt(s,9.08,5.67,3.75,1.05,'Soumitra Deshpande  ·  Prathamesh Swami\nPranav Kale  ·  R Virshin  ·  Raj Koli',10.5,'cream')

    # 02 — Problem / thesis.
    s=base(prs,2,'01 / THE QUESTION');heading(s,'Can chemistry predict taste?','A numerical quality score is estimated from laboratory measurements of red wine.')
    txt(s,.73,2.04,7.85,1.35,'11 measurements\n→ 1 quality score',32,'cream',False,SERIF)
    line(s,.73,3.65,7.97,3.65,'wine2',3)
    txt(s,.75,3.96,6.6,.55,'Supervised regression',20,'gold',True)
    txt(s,.75,4.56,6.35,1.1,'The model predicts the dataset’s rated quality value. We compare predictions with held-out ratings using error and fit metrics.',15,'muted')
    rect(s,9.0,2.03,3.69,3.97,'wine')
    txt(s,9.36,2.37,2.9,.3,'INPUT',10,'gold',True)
    txt(s,9.36,2.82,2.88,.84,'acidity · sulphates\nalcohol · density · …',16,'cream')
    line(s,9.38,3.91,12.25,3.91,'gold',1.2)
    txt(s,9.36,4.18,2.9,.3,'OUTPUT',10,'gold',True)
    txt(s,9.36,4.61,2.75,.65,'quality',27,'cream',False,SERIF)
    txt(s,9.37,5.26,2.67,.35,'observed labels 3–8',12,'cream')
    txt(s,.75,6.3,11.5,.29,'Goal: compare a single Decision Tree with a 300-tree Random Forest, then serve the better pipeline in Streamlit.',12.4,'cream')

    # 03 — Dataset / real distribution.
    s=base(prs,3,'02 / THE DATA');heading(s,'The UCI red-wine dataset','1,599 rows · 11 physicochemical inputs · 1 numeric target')
    stats=[('1,599','observations'),('11','numeric predictors'),('0','missing cells'),('240','duplicate rows')]
    for i,(v,l) in enumerate(stats):
        x=.66+i*3.1
        txt(s,x,2.05,2.7,.7,v,30,'gold',False,SERIF)
        txt(s,x,2.8,2.75,.28,l.upper(),10,'muted',True)
        if i<3: line(s,x+2.78,2.06,x+2.78,3.13,'line',.8)
    txt(s,.67,3.54,5.8,.35,'QUALITY SCORE DISTRIBUTION',12,'gold',True)
    counts=DATA['quality'].value_counts().sort_index()
    maxv=counts.max()
    for i,(score,count) in enumerate(counts.items()):
        x=.77+i*1.37; bh=2.15*count/maxv
        rect(s,x,6.14-bh,.77,bh,'gold' if score in (5,6) else 'wine2')
        txt(s,x,6.22,.77,.25,str(score),12,'cream',True,align=PP_ALIGN.CENTER)
        txt(s,x,5.85-bh,.77,.22,str(count),10,'muted',True,align=PP_ALIGN.CENTER)
    txt(s,9.3,3.58,3.18,.33,'WHY THIS MATTERS',11,'gold',True)
    txt(s,9.3,4.05,3.2,1.35,'Ratings 5 and 6 account for 82.5% of samples. Extreme scores are rare, so the model mostly learns the middle of the scale.',15,'cream')
    link(s,9.3,6.22,3.1,.25,'UCI source ↗','https://archive.ics.uci.edu/dataset/186/wine+quality',11)

    # 04 — EDA, actual correlations.
    s=base(prs,4,'03 / EXPLORATORY ANALYSIS');heading(s,'Signals in the raw data','Pearson correlation with quality highlights directional relationships before modeling.')
    corr=DATA.corr(numeric_only=True)['quality'].drop('quality')
    selected=[('Alcohol',corr['alcohol']),('Sulphates',corr['sulphates']),('Citric acid',corr['citric acid']),('Total sulfur dioxide',corr['total sulfur dioxide']),('Volatile acidity',corr['volatile acidity'])]
    txt(s,.78,2.03,4.5,.32,'CORRELATION WITH QUALITY',12,'gold',True)
    x0=5.71
    line(s,x0,2.59,x0,5.77,'line',1.1)
    for i,(name,v) in enumerate(selected):
        y=2.62+i*.65
        txt(s,.8,y,2.75,.26,name,14,'cream')
        line(s,x0,y+.13,x0+v*5.8,y+.13,'gold' if v>0 else 'wine2',8)
        txt(s,9.0,y,1.0,.27,f'{v:+.3f}',13,'gold' if v>0 else 'blush',True)
    rect(s,10.15,2.08,2.5,3.9,'panel')
    txt(s,10.43,2.4,1.98,.26,'TAKEAWAY',10,'gold',True)
    txt(s,10.43,2.95,1.95,2.48,'Alcohol has the strongest positive linear link; volatile acidity the strongest negative. These are associations, not causes.',14,'cream')
    txt(s,.78,6.23,10.9,.36,'The notebook also plots the full correlation matrix and checks the distribution of the target labels.',12,'muted')

    # 05 — preprocessing.
    s=base(prs,5,'04 / PREPARATION');heading(s,'A consistent path to model-ready data','The notebook and training script use the same scikit-learn workflow.')
    steps=[('01','READ','Semicolon-separated CSV'),('02','SPLIT','80% train / 20% test'),('03','IMPUTE','Median per feature'),('04','SCALE','StandardScaler'),('05','FIT','Tree or forest')]
    for i,(n,t,d) in enumerate(steps):
        x=.67+i*2.52
        rect(s,x,2.23,2.19,2.42,'wine' if i==4 else 'panel')
        txt(s,x+.2,2.42,1.5,.38,n,24,'gold',False,SERIF)
        txt(s,x+.2,3.02,1.8,.29,t,13,'cream',True)
        txt(s,x+.2,3.52,1.8,.69,d,12,'muted' if i<4 else 'cream')
        if i<4: line(s,x+2.2,3.37,x+2.5,3.37,'gold',1.5)
    line(s,.73,5.32,12.62,5.32,'line',.8)
    txt(s,.74,5.64,4.05,.29,'REPRODUCIBLE SPLIT',11,'gold',True)
    txt(s,.74,6.04,4.06,.44,'random_state=42  ·  1,279 training rows  ·  320 test rows',12,'cream')
    txt(s,6.78,5.64,4.6,.29,'PIPELINE SAFETY',11,'gold',True)
    txt(s,6.78,6.04,5.62,.5,'Imputer and scaler are fitted with training data, then reused for test data and app predictions.',12,'cream')

    # 06 — decision tree.
    s=base(prs,6,'05 / MODEL 01');heading(s,'Decision Tree Regressor','A single constrained tree creates interpretable if/then regions.')
    txt(s,.75,2.06,5.2,.36,'ONE MODEL · ONE PATH PER WINE',12,'gold',True)
    rect(s,.77,2.69,4.58,.7,'wine')
    txt(s,1.0,2.88,4.1,.3,'Root: choose the strongest split',15,'cream',True)
    line(s,3.06,3.39,1.83,4.1,'gold',1.7);line(s,3.06,3.39,4.31,4.1,'gold',1.7)
    rect(s,.87,4.1,2.03,.7,'panel');rect(s,3.3,4.1,2.03,.7,'panel')
    txt(s,1.02,4.31,1.72,.25,'feature ≤ threshold',11,'cream');txt(s,3.45,4.31,1.72,.25,'feature > threshold',11,'cream')
    line(s,1.88,4.8,1.88,5.43,'gold',1.5);line(s,4.32,4.8,4.32,5.43,'gold',1.5)
    oval(s,1.58,5.43,.6,.6,'gold');oval(s,4.02,5.43,.6,.6,'gold')
    txt(s,.77,6.28,5.3,.26,'Terminal leaves return a mean quality estimate.',12,'muted')
    rect(s,6.9,2.02,5.75,4.68,'panel')
    txt(s,7.25,2.37,4.9,.29,'CONFIGURATION',11,'gold',True)
    txt(s,7.25,2.86,4.7,.34,'max_depth = 5',17,'cream',True)
    txt(s,7.25,3.34,4.7,.34,'min_samples_leaf = 3',17,'cream',True)
    line(s,7.25,3.93,12.29,3.93,'line',.8)
    txt(s,7.25,4.2,2.0,.27,'HELD-OUT TEST',10,'muted',True)
    txt(s,7.25,4.63,2.29,.59,'0.660239',26,'cream',False,SERIF)
    txt(s,7.25,5.29,2.3,.25,'RMSE ↓',11,'gold',True)
    txt(s,10.1,4.63,2.2,.59,'0.332959',26,'cream',False,SERIF)
    txt(s,10.1,5.29,2.2,.25,'R² ↑',11,'gold',True)
    txt(s,7.25,6.02,4.81,.3,'Useful baseline, but one tree has higher test error.',11,'muted')

    # 07 — Random forest.
    s=base(prs,7,'06 / MODEL 02');heading(s,'Random Forest Regressor','Three hundred randomized trees average their predictions into one estimate.')
    for i in range(7):
        x=.78+i*.72
        line(s,x+.2,3.03,x+.2,4.55,'gold',1)
        line(s,x+.2,3.52,x-.08,3.86,'gold',1)
        line(s,x+.2,3.52,x+.48,3.86,'gold',1)
        line(s,x+.2,4.17,x-.02,4.43,'gold',1)
        line(s,x+.2,4.17,x+.42,4.43,'gold',1)
        oval(s,x+.09,2.76,.22,.22,'wine2')
    txt(s,.87,2.19,5.2,.3,'ENSEMBLE LOGIC',12,'gold',True)
    line(s,1.14,5.2,5.52,5.2,'gold',1.7)
    txt(s,1.02,5.55,5.25,.45,'tree₁ + tree₂ + … + tree₃₀₀',19,'cream',False,SERIF)
    txt(s,2.43,6.11,2.95,.34,'÷ 300  →  prediction',14,'gold',True)
    rect(s,6.9,2.02,5.75,4.68,'wine')
    txt(s,7.25,2.37,4.9,.29,'SELECTED MODEL',11,'gold',True)
    txt(s,7.25,2.86,4.7,.34,'n_estimators = 300',17,'cream',True)
    txt(s,7.25,3.34,4.7,.34,'min_samples_leaf = 2',17,'cream',True)
    line(s,7.25,3.93,12.29,3.93,'gold',.8)
    txt(s,7.25,4.2,2.0,.27,'HELD-OUT TEST',10,'gold',True)
    txt(s,7.25,4.63,2.29,.59,'0.557963',26,'cream',False,SERIF)
    txt(s,7.25,5.29,2.3,.25,'RMSE ↓',11,'gold',True)
    txt(s,10.1,4.63,2.2,.59,'0.523612',26,'cream',False,SERIF)
    txt(s,10.1,5.29,2.2,.25,'R² ↑',11,'gold',True)
    txt(s,7.25,6.02,4.81,.32,'Lower RMSE wins the notebook’s selection rule.',11,'cream')

    # 08 — implementation.
    s=base(prs,8,'07 / IMPLEMENTATION');heading(s,'One workflow, saved as a reusable pipeline','The notebook documents exploration and evaluation; train_model.py exports the fitted artifacts.')
    phases=[('01','NOTEBOOK','Inspect data, plot target distribution and correlations.'),('02','TRAIN','Fit both pipelines on identical training rows.'),('03','COMPARE','Calculate MAE, MSE, RMSE and R²; select by RMSE.'),('04','SAVE','Joblib model, metrics, predictions and feature importance.')]
    for i,(n,t,d) in enumerate(phases):
        x=.69+i*3.13
        rect(s,x,2.1,2.83,3.55,'panel' if i%2==0 else 'bg2')
        rect(s,x,2.1,2.83,.055,'gold' if i==3 else 'wine2')
        txt(s,x+.22,2.46,2.36,.43,n,23,'gold',False,SERIF)
        txt(s,x+.22,3.14,2.4,.29,t,13,'cream',True)
        txt(s,x+.22,3.67,2.32,1.18,d,13,'muted')
    rect(s,.69,5.95,11.96,.73,'wine')
    txt(s,.96,6.18,11.5,.29,'Saved artifact:  models/wine_quality_model.joblib   =   median imputer  +  scaler  +  selected forest',13,'cream',True)

    # 09 — exact metrics + editorial comparison.
    s=base(prs,9,'08 / TEST RESULTS');heading(s,'The forest improves every test metric','Both estimators used the same 320 held-out rows. Lower errors and higher R² are better.')
    headers=[('MODEL',.77,2.13,3.18),('MAE',4.34,2.13,1.52),('MSE',6.02,2.13,1.52),('RMSE',7.7,2.13,1.67),('R²',9.66,2.13,1.52)]
    for name,x,y,w in headers: txt(s,x,y,w,.24,name,11,'gold',True)
    line(s,.73,2.51,12.42,2.51,'gold',1.2)
    rows=[('Decision Tree',MET['Decision Tree Regressor'],'panel'),('Random Forest',MET['Random Forest Regressor'],'wine')]
    for i,(name,m,fill) in enumerate(rows):
        y=2.69+i*.91;rect(s,.73,y,11.82,.74,fill)
        txt(s,.93,y+.2,3.1,.31,name,15,'cream',True)
        for key,x,w in [('MAE',4.34,1.52),('MSE',6.02,1.52),('RMSE',7.7,1.67),('R2',9.66,1.52)]:
            txt(s,x,y+.2,w,.31,f'{m[key]:.6f}',15,'cream',i==1)
    txt(s,.77,4.88,5.42,.34,'RMSE  ·  quality-score units',12,'gold',True)
    dt=MET['Decision Tree Regressor']['RMSE'];rf=MET['Random Forest Regressor']['RMSE']
    bar(s,.8,5.47,5.23,.29,dt,.73,'blush')
    bar(s,.8,5.91,5.23,.29,rf,.73,'gold')
    txt(s,6.2,5.42,2.2,.34,'Tree    0.660',14,'cream')
    txt(s,6.2,5.88,2.2,.34,'Forest  0.558',14,'cream',True)
    rect(s,9.17,4.88,3.37,1.45,'panel')
    txt(s,9.46,5.09,2.85,.25,'RMSE IMPROVEMENT',10,'gold',True)
    txt(s,9.45,5.46,2.8,.52,f'{(1-rf/dt)*100:.2f}%',26,'cream',False,SERIF)
    txt(s,.77,6.55,11.5,.25,'Exact six-decimal values are from artifacts/metrics.json, generated by the notebook/training workflow.',10,'muted')

    # 10 — importance.
    s=base(prs,10,'09 / INTERPRETATION');heading(s,'What the forest relied on','Impurity-based feature importance measures use in fitted tree splits, not causal effects.')
    vals=FI.head(8)
    m=FI.importance.max()
    for i,row in enumerate(vals.itertuples()):
        y=2.04+i*.53
        txt(s,.79,y,3.18,.27,row.feature.title(),13,'cream',i<3)
        bar(s,4.01,y+.04,6.37,.19,row.importance,m,'gold' if i==0 else 'wine2','panel')
        txt(s,10.62,y,1.34,.27,f'{row.importance*100:.1f}%',13,'gold' if i==0 else 'cream',True)
    rect(s,.79,6.42,11.75,.39,'panel')
    txt(s,1.0,6.5,11.25,.19,'Top three: alcohol 28.4%  ·  sulphates 15.9%  ·  volatile acidity 11.0%',11,'cream',True)

    # 11 — faithful Streamlit browser UI mock; source values are median profile prediction.
    s=base(prs,11,'10 / DEPLOYMENT');heading(s,'The model becomes a usable app','Streamlit loads the saved pipeline and accepts the same 11 feature names and training ranges.')
    rect(s,.7,2.02,11.93,4.75,'white')
    rect(s,.7,2.02,11.93,.38,'#E7E4E6')
    oval(s,.91,2.16,.1,.1,'#E16A70');oval(s,1.09,2.16,.1,.1,'#E7B75E');oval(s,1.27,2.16,.1,.1,'#70B881')
    txt(s,4.42,2.12,4.25,.15,'Wine Quality Predictor — Streamlit',8.5,'#5F5B60',True,align=PP_ALIGN.CENTER)
    rect(s,.7,2.4,2.27,4.37,'#F3F1F4')
    txt(s,.93,2.76,1.8,.27,'About this project',10,'#362B36',True)
    txt(s,.93,3.12,1.77,.68,'A model estimates red-wine quality from laboratory measurements.',8.5,'#5D5660')
    line(s,.93,3.96,2.71,3.96,'#D8D3D9',.8)
    txt(s,.93,4.23,1.76,.2,'Selected model',8,'#776E78',True)
    txt(s,.93,4.53,1.8,.4,'Random Forest\nRegressor',9.2,'#362B36')
    txt(s,.93,5.29,1.76,.2,'Dataset',8,'#776E78',True)
    txt(s,.93,5.59,1.8,.36,'UCI Red Wine Quality',9,'#362B36')
    rect(s,3.2,2.64,9.08,1.04,'wine')
    txt(s,3.44,2.8,7.7,.21,'MACHINE LEARNING FUNDAMENTALS',8.5,'gold',True)
    txt(s,3.44,3.12,7.9,.45,'Wine Quality Predictor',22,'white',True)
    txt(s,3.38,3.92,3.4,.23,'Predict quality',10,'#50243B',True)
    txt(s,6.24,3.92,2.4,.23,'Model insights',10,'#7B747B')
    txt(s,8.9,3.92,2.2,.23,'Project guide',10,'#7B747B')
    line(s,3.39,4.23,12.21,4.23,'#DDDADE',.8)
    txt(s,3.42,4.46,4.3,.28,'Enter wine properties',14,'#26222A',True)
    # Narrow native input controls, consistent with three-column form.
    input_labels=[('Fixed Acidity','7.9000'),('Volatile Acidity','0.5200'),('Citric Acid','0.2600'),('Residual Sugar','2.2000'),('Chlorides','0.0790'),('Alcohol','10.2000')]
    for i,(lab,val) in enumerate(input_labels):
        col=i%3; row=i//3; x=3.43+col*1.85;y=4.93+row*.68
        txt(s,x,y,1.72,.18,lab,7.4,'#504B51',True)
        rect(s,x,y+.22,1.58,.32,'white','#CFC9D0')
        txt(s,x+.08,y+.29,1.38,.16,val,8.2,'#342F36')
    rect(s,3.43,6.36,2.61,.29,'wine')
    txt(s,3.55,6.42,2.36,.15,'Predict wine quality',8.5,'white',True,align=PP_ALIGN.CENTER)
    txt(s,6.25,6.41,2.95,.18,'+ 5 more inputs below the fold',8.2,'#6D6670',True)
    rect(s,9.48,4.55,2.78,1.87,'#FFF5F8','#EED8E2')
    txt(s,9.7,4.79,2.3,.2,'PREDICTED QUALITY SCORE',8,'#8F5870',True)
    txt(s,9.7,5.18,2.3,.62,'5.69 / 10',25,'#5A1740',True)
    txt(s,9.7,5.93,2.26,.2,'Middle predicted range',8.5,'#6D214F',True)
    txt(s,.84,6.82,11.8,.17,'UI mock reflects app.py; 5.69 is the saved model’s prediction for the median-value profile.',9,'muted')

    # 12 — limitations, concrete next steps.
    s=base(prs,12,'11 / LIMITATIONS');heading(s,'How far should we trust it?','A useful educational estimate still has boundaries.')
    limits=[('01','Subjective target','Quality scores reflect human ratings, so labels contain judgment.'),('02','Uneven labels','Most examples are scores 5–6; extremes have little training support.'),('03','Single split + duplicates','The reported test metrics use one split of data containing 240 duplicate rows.')]
    for i,(n,t,d) in enumerate(limits):
        y=2.03+i*1.44
        txt(s,.75,y,.62,.37,n,19,'gold',False,SERIF)
        txt(s,1.56,y,4.95,.34,t,18,'cream',True)
        txt(s,1.57,y+.44,4.97,.6,d,13,'muted')
        line(s,.75,y+1.18,6.52,y+1.18,'line',.7)
    rect(s,7.24,2.02,5.37,4.5,'wine')
    txt(s,7.62,2.42,4.45,.29,'NEXT ITERATION',11,'gold',True)
    txt(s,7.62,2.95,4.47,.65,'Test robustness, then broaden the data.',20,'cream',False,SERIF)
    txt(s,7.62,4.11,4.44,1.74,'Group duplicate wines before splitting.\nUse repeated cross-validation.\nEvaluate performance at rare scores.\nAdd sensory and contextual features.',14,'cream')
    txt(s,.76,6.71,10.8,.2,'Use the prediction as a model-based estimate, not an official wine rating.',11,'gold',True)

    # 13 — conclusion and team / access.
    s=base(prs,13,'12 / CONCLUSION');heading(s,'From data to decision','A complete machine-learning workflow with a working inference interface.')
    rect(s,.7,2.04,6.53,4.54,'wine')
    txt(s,1.04,2.4,5.75,.29,'SELECTED PIPELINE',11,'gold',True)
    txt(s,1.04,2.92,5.72,.57,'Random Forest',27,'cream',False,SERIF)
    txt(s,1.04,3.62,5.72,.39,'RMSE 0.557963   ·   R² 0.523612',17,'cream',True)
    line(s,1.04,4.3,6.74,4.3,'gold',1.1)
    txt(s,1.04,4.64,5.54,.64,'The fitted preprocessing and forest are saved together and used by Streamlit.',14,'cream')
    txt(s,1.04,5.81,5.43,.33,'TEAM BDA  /  GROUP 14',12,'gold',True)
    txt(s,7.7,2.09,4.87,.3,'PROJECT LINKS',12,'gold',True)
    link(s,7.7,2.53,4.8,.61,'GitHub repository ↗','https://github.com/SoumitraDeshpande11/wine-quality-prediction',18)
    txt(s,7.7,2.99,4.82,.43,'github.com/SoumitraDeshpande11/\nwine-quality-prediction',10,'muted')
    line(s,7.7,3.6,12.46,3.6,'line',.8)
    link(s,7.7,3.83,4.8,.55,'Live Streamlit app ↗','https://soumitradeshpande11-wine-quality-prediction-app-pnq8pq.streamlit.app/',18)
    txt(s,7.7,4.29,4.82,.54,'soumitradeshpande11-wine-quality-\nprediction-app-pnq8pq.streamlit.app',10,'muted')
    line(s,7.7,4.98,12.46,4.98,'line',.8)
    txt(s,7.7,5.21,4.84,.26,'PRESENTED BY',11,'gold',True)
    txt(s,7.7,5.57,4.83,1.1,'Soumitra Deshpande · Prathamesh Swami\nPranav Kale · R Virshin · Raj Koli',12,'cream')

    prs.core_properties.title='Wine Quality Prediction | Team BDA'
    prs.core_properties.subject='Red-wine regression model comparison and Streamlit deployment'
    prs.core_properties.author='Team BDA'
    prs.save(OUT)
    print(OUT)
if __name__=='__main__':make()
