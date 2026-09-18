from pathlib import Path
import json
import math
import textwrap
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
from matplotlib.colors import LinearSegmentedColormap
from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_THEME_COLOR

ROOT = Path(__file__).resolve().parent.parent
ASSET = ROOT / "presentation" / "assets"
ASSET.mkdir(parents=True, exist_ok=True)
OUT = ROOT / "presentation" / "Wine_Quality_Prediction_BDA.pptx"

W, H = 13.333, 7.5
BG = "17161B"
PANEL = "242229"
PANEL2 = "2C2932"
RED = "8F2F45"
RED2 = "B9495B"
GOLD = "E6B86A"
GOLD2 = "F2D39A"
INK = "F7F2EA"
MUTED = "BDB7AE"
MUTED2 = "918B85"
GREEN = "7ECFA2"
BLUE = "8FB9E8"
BLACK = "111014"
FONT = "Aptos"
DISPLAY = "Aptos Display"


def rgb(hexstr):
    hexstr = hexstr.replace("#", "")
    return RGBColor(int(hexstr[0:2], 16), int(hexstr[2:4], 16), int(hexstr[4:6], 16))


def hexrgb(hexstr):
    return tuple(int(hexstr[i:i+2], 16) for i in (0, 2, 4))


def make_gradient(path, accent=(143, 47, 69), accent2=(230, 184, 106)):
    width, height = 1600, 900
    y, x = np.mgrid[0:height, 0:width]
    base = np.zeros((height, width, 3), dtype=float)
    b = np.array(hexrgb(BG), dtype=float)
    base[:] = b
    glow1 = np.exp(-(((x - 1260) / 430) ** 2 + ((y - 80) / 370) ** 2))[:, :, None]
    glow2 = np.exp(-(((x - 160) / 470) ** 2 + ((y - 790) / 520) ** 2))[:, :, None]
    base = base * 0.85 + np.array(accent, dtype=float) * glow1 * 0.44 + np.array(accent2, dtype=float) * glow2 * 0.08
    base = np.clip(base, 0, 255).astype(np.uint8)
    image = Image.fromarray(base, "RGB")
    draw = ImageDraw.Draw(image, "RGBA")
    for i in range(5):
        box = (1100 - i * 42, -180 + i * 38, 1800 - i * 42, 520 + i * 38)
        draw.arc(box, 205, 330, fill=(*accent2, 44 - i * 6), width=3)
    for i in range(3):
        draw.line((0, 745 + i * 18, 980, 675 + i * 18), fill=(*accent, 30), width=2)
    image.save(path)


def make_chart_assets():
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10})
    dark = "#" + BG
    white = "#" + INK
    gold = "#" + GOLD
    red = "#" + RED2
    muted = "#" + MUTED
    data = pd.read_csv(ROOT / "data" / "winequality-red.csv", sep=";")
    metrics = json.loads((ROOT / "artifacts" / "metrics.json").read_text())
    fi = pd.read_csv(ROOT / "artifacts" / "feature_importance.csv").sort_values("importance")

    fig, ax = plt.subplots(figsize=(7.2, 3.4), dpi=220)
    fig.patch.set_facecolor(dark); ax.set_facecolor(dark)
    counts = data["quality"].value_counts().sort_index()
    bars = ax.bar(counts.index.astype(str), counts.values, color=[red if v not in [5,6] else gold for v in counts.index], width=0.64)
    ax.set_title("Quality labels are concentrated around 5–6", color=white, loc="left", fontsize=14, fontweight="bold", pad=14)
    ax.set_xlabel("Quality score", color=muted, labelpad=8); ax.set_ylabel("Number of wines", color=muted, labelpad=8)
    ax.tick_params(colors=muted); ax.grid(axis="y", color="#49454f", alpha=0.35, linewidth=0.8)
    for spine in ax.spines.values(): spine.set_visible(False)
    for bar in bars:
        ax.text(bar.get_x()+bar.get_width()/2, bar.get_height()+15, f"{int(bar.get_height())}", color=white, ha="center", fontsize=9)
    fig.tight_layout()
    fig.savefig(ASSET / "quality_distribution_dark.png", facecolor=dark, bbox_inches="tight", transparent=False)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(7.2, 3.6), dpi=220)
    fig.patch.set_facecolor(dark); ax.set_facecolor(dark)
    labels = [x.replace("total ", "total ").replace("free ", "free ") for x in fi["feature"]]
    vals = fi["importance"] * 100
    cols = [gold if v == vals.max() else red2 if False else red for v in vals]
    cols[-1] = gold
    ax.barh(labels, vals, color=cols, height=0.64)
    ax.set_title("Random Forest feature importance", color=white, loc="left", fontsize=14, fontweight="bold", pad=14)
    ax.set_xlabel("Relative importance (%)", color=muted, labelpad=8)
    ax.tick_params(axis="x", colors=muted); ax.tick_params(axis="y", colors=white, labelsize=9)
    ax.grid(axis="x", color="#49454f", alpha=0.35, linewidth=0.8)
    for spine in ax.spines.values(): spine.set_visible(False)
    for y0, v in enumerate(vals): ax.text(v+0.8, y0, f"{v:.1f}%", color=white, va="center", fontsize=8.5)
    fig.tight_layout()
    fig.savefig(ASSET / "feature_importance_dark.png", facecolor=dark, bbox_inches="tight", transparent=False)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(6.7, 3.45), dpi=220)
    fig.patch.set_facecolor(dark); ax.set_facecolor(dark)
    models = ["Decision Tree", "Random Forest"]
    x = np.arange(2); width = 0.32
    rmse = [metrics["models"]["Decision Tree Regressor"]["RMSE"], metrics["models"]["Random Forest Regressor"]["RMSE"]]
    r2 = [metrics["models"]["Decision Tree Regressor"]["R2"], metrics["models"]["Random Forest Regressor"]["R2"]]
    bars1 = ax.bar(x-width/2, rmse, width, label="RMSE ↓", color=red)
    bars2 = ax.bar(x+width/2, r2, width, label="R² ↑", color=gold)
    ax.set_ylim(0, 0.76); ax.set_xticks(x, models); ax.tick_params(colors=muted)
    ax.set_ylabel("Score", color=muted); ax.grid(axis="y", color="#49454f", alpha=0.35, linewidth=0.8)
    ax.set_title("Random Forest delivers the stronger test-set fit", color=white, loc="left", fontsize=14, fontweight="bold", pad=14)
    for spine in ax.spines.values(): spine.set_visible(False)
    for bars in (bars1, bars2):
        for bar in bars: ax.text(bar.get_x()+bar.get_width()/2, bar.get_height()+0.025, f"{bar.get_height():.3f}", color=white, ha="center", fontsize=9)
    leg = ax.legend(frameon=False, loc="upper left", ncol=2, bbox_to_anchor=(0, 1.02));
    for t in leg.get_texts(): t.set_color(white)
    fig.tight_layout()
    fig.savefig(ASSET / "model_comparison_dark.png", facecolor=dark, bbox_inches="tight", transparent=False)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(6.8, 3.6), dpi=220)
    fig.patch.set_facecolor(dark); ax.set_facecolor(dark)
    corr = data.corr(numeric_only=True)
    cmap = LinearSegmentedColormap.from_list("wine", ["#292630", "#6f273d", "#e6b86a"])
    im = ax.imshow(corr.values, cmap=cmap, vmin=-1, vmax=1, aspect="auto")
    labels = ["fixed\nacidity", "volatile\nacidity", "citric\nacid", "residual\nsugar", "chlorides", "free SO₂", "total SO₂", "density", "pH", "sulphates", "alcohol", "quality"]
    ax.set_xticks(np.arange(len(labels)), labels, rotation=52, ha="right", fontsize=7, color=muted); ax.set_yticks(np.arange(len(labels)), labels, fontsize=7, color=muted)
    ax.set_title("Feature relationships to understand before modeling", color=white, loc="left", fontsize=13, fontweight="bold", pad=12)
    for spine in ax.spines.values(): spine.set_visible(False)
    fig.colorbar(im, ax=ax, fraction=0.025, pad=0.02).ax.tick_params(colors=muted, labelsize=7)
    fig.tight_layout()
    fig.savefig(ASSET / "correlation_dark.png", facecolor=dark, bbox_inches="tight", transparent=False)
    plt.close(fig)


def add_bg(slide, image=None):
    if image:
        slide.shapes.add_picture(str(image), 0, 0, width=Inches(W), height=Inches(H))
    else:
        fill = slide.background.fill; fill.solid(); fill.fore_color.rgb = rgb(BG)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.26), Inches(W), Inches(0.24))
    band.fill.solid(); band.fill.fore_color.rgb = rgb(RED); band.line.fill.background()


def box(slide, x, y, w, h, fill=PANEL, line=None, radius=False, transparency=0):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill); shape.fill.transparency = transparency
    if line:
        shape.line.color.rgb = rgb(line); shape.line.width = Pt(0.7)
    else:
        shape.line.fill.background()
    return shape


def text(slide, x, y, w, h, value, size=16, color=INK, bold=False, font=FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.03, italic=False, caps=False):
    tx = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tx.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(margin); tf.margin_right = Inches(margin); tf.margin_top = Inches(margin); tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = value.upper() if caps else value
    run.font.name = font; run.font.size = Pt(size); run.font.bold = bold; run.font.italic = italic; run.font.color.rgb = rgb(color)
    return tx


def rich_text(slide, x, y, w, h, runs, size=16, color=INK, font=FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.03, spacing=1.0):
    tx = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=tx.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=Inches(margin); tf.margin_right=Inches(margin); tf.margin_top=Inches(margin); tf.margin_bottom=Inches(margin); tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.alignment=align
    for run_spec in runs:
        r=p.add_run(); r.text=run_spec[0]; r.font.name=run_spec[3] if len(run_spec)>3 and run_spec[3] else font; r.font.size=Pt(run_spec[1] if len(run_spec)>1 and run_spec[1] else size); r.font.bold=run_spec[2] if len(run_spec)>2 else False; r.font.color.rgb=rgb(run_spec[4] if len(run_spec)>4 and run_spec[4] else color)
    p.line_spacing = spacing
    return tx


def label(slide, x, y, value, color=GOLD, size=9.5):
    text(slide, x, y, 2.5, 0.22, value, size=size, color=color, bold=True, caps=True)


def title(slide, kicker, heading, sub=None, page=None):
    label(slide, 0.62, 0.39, kicker)
    text(slide, 0.62, 0.69, 12.1, 0.62, heading, size=28, color=INK, bold=True, font=DISPLAY)
    if sub: text(slide, 0.64, 1.34, 11.8, 0.35, sub, size=11.5, color=MUTED)
    if page is not None:
        text(slide, 12.58, 7.31, 0.54, 0.14, f"{page:02d}", size=8, color=INK, bold=True, align=PP_ALIGN.RIGHT, valign=MSO_ANCHOR.MIDDLE, margin=0)


def rule(slide, x, y, w, color=RED2, height=0.035):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(height)); sh.fill.solid(); sh.fill.fore_color.rgb=rgb(color); sh.line.fill.background(); return sh


def pill(slide, x, y, w, value, fill=PANEL2, color=MUTED, size=9.5):
    box(slide,x,y,w,0.31,fill=fill,radius=True)
    text(slide,x+0.08,y+0.055,w-0.16,0.19,value,size=size,color=color,bold=True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE,margin=0)


def add_bullet(slide, x, y, w, value, accent=GOLD, size=13, h=0.39):
    dot=slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y+0.11), Inches(0.10), Inches(0.10)); dot.fill.solid(); dot.fill.fore_color.rgb=rgb(accent); dot.line.fill.background()
    text(slide,x+0.20,y,w-0.20,h,value,size=size,color=INK)


def add_link(slide, x,y,w,h,label_text,url, color=GOLD, size=10):
    tx=text(slide,x,y,w,h,label_text,size=size,color=color,bold=True)
    try:
        tx.text_frame.paragraphs[0].runs[0].hyperlink.address=url
    except Exception:
        pass
    return tx


def codebox(slide, x, y, w, h, lines, highlight=None, size=10.5):
    box(slide,x,y,w,h,fill=BLACK,line="#3B3640",radius=True)
    tx=slide.shapes.add_textbox(Inches(x+0.18),Inches(y+0.13),Inches(w-0.36),Inches(h-0.24)); tf=tx.text_frame; tf.clear(); tf.word_wrap=False
    tf.margin_left=Inches(0); tf.margin_right=Inches(0); tf.margin_top=Inches(0); tf.margin_bottom=Inches(0)
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.space_after=Pt(2); r=p.add_run(); r.text=line; r.font.name="Menlo"; r.font.size=Pt(size); r.font.color.rgb=rgb(GOLD2 if highlight and i in highlight else INK)
    return tx


def metric_card(slide,x,y,w,label_text,value,sub,color=GOLD):
    box(slide,x,y,w,1.05,fill=PANEL,line="#39343E",radius=True)
    label(slide,x+0.16,y+0.14,label_text,color=color,size=8.5)
    text(slide,x+0.16,y+0.39,w-0.32,0.32,value,size=24,color=INK,bold=True,font=DISPLAY)
    text(slide,x+0.16,y+0.76,w-0.32,0.17,sub,size=8.5,color=MUTED)


def workflow_node(slide,x,y,w,h,number,heading,body,accent=GOLD):
    box(slide,x,y,w,h,fill=PANEL,line="#3B3640",radius=True)
    circ=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x+0.18),Inches(y+0.18),Inches(0.35),Inches(0.35)); circ.fill.solid();circ.fill.fore_color.rgb=rgb(accent);circ.line.fill.background()
    text(slide,x+0.18,y+0.245,0.35,0.13,str(number),size=9,color=BG,bold=True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE,margin=0)
    text(slide,x+0.66,y+0.19,w-0.84,0.22,heading,size=11,color=INK,bold=True)
    text(slide,x+0.18,y+0.65,w-0.36,h-0.76,body,size=9.2,color=MUTED)


def add_arrow(slide,x,y,w=0.34,color=GOLD):
    a=slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(0.18)); a.fill.solid(); a.fill.fore_color.rgb=rgb(color); a.line.fill.background(); return a


def build():
    make_gradient(ASSET / "cover_bg.png")
    make_chart_assets()
    data = pd.read_csv(ROOT / "data" / "winequality-red.csv", sep=";")
    metrics = json.loads((ROOT / "artifacts" / "metrics.json").read_text())
    fi = pd.read_csv(ROOT / "artifacts" / "feature_importance.csv").sort_values("importance", ascending=False)

    prs=Presentation(); prs.slide_width=Inches(W); prs.slide_height=Inches(H)
    blank=prs.slide_layouts[6]

    s=prs.slides.add_slide(blank); add_bg(s,ASSET/"cover_bg.png")
    box(s,0.68,0.68,1.22,0.31,fill=GOLD,radius=True)
    text(s,0.78,0.735,1.02,0.16,"GROUP 14",size=9,color=BG,bold=True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE,margin=0)
    text(s,0.72,1.53,7.35,1.30,"Predicting\nWine Quality",size=34,color=INK,bold=True,font=DISPLAY)
    text(s,0.76,3.07,6.65,0.46,"A machine learning regression system built from red-wine chemistry",size=15,color=GOLD2)
    rule(s,0.76,3.73,1.08,GOLD,height=0.055)
    text(s,0.76,4.04,6.60,0.28,"TEAM BDA",size=12,color=GOLD,bold=True)
    text(s,0.76,4.43,6.40,0.74,"Soumitra Deshpande  ·  Prathamesh Swami  ·  Pranav Kale\nR Virshin  ·  Raj Koli",size=11.5,color=INK)
    box(s,8.43,1.47,3.95,3.88,fill=BLACK,line="#3B3640",radius=True,transparency=8)
    label(s,8.78,1.84,"PROJECT SNAPSHOT",color=GOLD)
    metric_card(s,8.78,2.20,1.56,"SAMPLES","1,599","red wines",color=RED2)
    metric_card(s,10.52,2.20,1.56,"FEATURES","11","chemistry inputs",color=GOLD)
    metric_card(s,8.78,3.50,3.30,"FINAL MODEL","R² 0.524","Random Forest · test set",color=GOLD)
    pill(s,8.80,4.84,1.36,"REGRESSION",fill=RED,color=INK)
    pill(s,10.28,4.84,1.64,"STREAMLIT",fill=PANEL2,color=GOLD2)
    text(s,0.77,6.65,9.2,0.22,"ML FUNDAMENTALS  ·  MODEL COMPARISON  ·  DEPLOYMENT",size=9,color=MUTED,caps=True)
    text(s,12.18,7.31,0.54,0.14,"01",size=8,color=INK,bold=True,align=PP_ALIGN.RIGHT,margin=0)

    s=prs.slides.add_slide(blank); add_bg(s); title(s,"01 · Framing the problem","From chemistry measurements to a quality score","A numerical prediction problem with a clear, testable target.",2)
    box(s,0.65,2.03,7.08,3.74,fill=PANEL,line="#39343E",radius=True)
    label(s,0.94,2.32,"PROBLEM STATEMENT",color=RED2)
    text(s,0.94,2.72,6.38,1.02,"Estimate a red wine’s quality score from its physicochemical properties.",size=22,color=INK,bold=True,font=DISPLAY)
    text(s,0.94,4.00,6.20,0.84,"The target is the rated quality value. Because the output is numeric, we evaluate the system as a regression model rather than using classification accuracy.",size=12.5,color=MUTED)
    pill(s,0.94,5.18,1.22,"TARGET",fill=RED,color=INK)
    pill(s,2.29,5.18,1.54,"QUALITY 3–8",fill=PANEL2,color=GOLD2)
    pill(s,3.97,5.18,1.68,"REGRESSION",fill=PANEL2,color=GOLD2)
    box(s,8.02,2.03,4.66,3.74,fill=PANEL,line="#39343E",radius=True)
    label(s,8.34,2.32,"PROJECT OBJECTIVES",color=GOLD)
    add_bullet(s,8.34,2.79,3.90,"Prepare and understand the public dataset")
    add_bullet(s,8.34,3.38,3.90,"Compare two suitable regression models",accent=RED2)
    add_bullet(s,8.34,3.97,3.90,"Evaluate with MAE, MSE, RMSE and R²")
    add_bullet(s,8.34,4.56,3.90,"Explain influential chemistry features",accent=RED2)
    add_bullet(s,8.34,5.15,3.90,"Deploy the selected pipeline in Streamlit")

    s=prs.slides.add_slide(blank); add_bg(s); title(s,"02 · Data","A compact, credible dataset for a complete ML lifecycle","The UCI red-wine dataset is small enough to understand and rich enough to compare models.",3)
    metric_card(s,0.68,1.96,1.78,"ROWS","1,599","observations",color=GOLD)
    metric_card(s,2.66,1.96,1.78,"INPUTS","11","numeric features",color=RED2)
    metric_card(s,4.64,1.96,1.78,"TARGET","3–8","quality labels",color=GOLD)
    metric_card(s,6.62,1.96,1.78,"SPLIT","80 / 20","train / test",color=RED2)
    box(s,0.68,3.34,5.88,3.14,fill=PANEL,line="#39343E",radius=True)
    label(s,0.97,3.62,"FEATURE SET",color=GOLD)
    features = ["fixed acidity","volatile acidity","citric acid","residual sugar","chlorides","free sulfur dioxide","total sulfur dioxide","density","pH","sulphates","alcohol"]
    for i, f in enumerate(features):
        col=i//6; row=i%6
        text(s,0.98+col*2.75,4.06+row*0.36,2.5,0.20,f,size=10.4,color=INK)
        rule(s,0.98+col*2.75,4.31+row*0.36,0.18,GOLD if i<6 else RED2,height=0.025)
    text(s,0.98,6.14,5.0,0.22,"All predictors are numeric; no categorical encoding is needed.",size=9.5,color=MUTED,italic=True)
    box(s,6.90,3.34,5.77,3.14,fill=PANEL,line="#39343E",radius=True)
    label(s,7.19,3.62,"TARGET DISTRIBUTION",color=RED2)
    s.shapes.add_picture(str(ASSET/"quality_distribution_dark.png"), Inches(7.08), Inches(3.96), width=Inches(5.38), height=Inches(2.25))
    text(s,7.19,6.16,4.90,0.20,"Most ratings sit at 5 and 6, with fewer extremes.",size=9.4,color=MUTED)
    add_link(s,10.30,0.39,2.25,0.20,"UCI DATASET ↗","https://archive.ics.uci.edu/dataset/186/wine+quality",color=GOLD,size=8.5)

    s=prs.slides.add_slide(blank); add_bg(s); title(s,"03 · Preparation","One reproducible path from CSV to model-ready data","Preprocessing is kept inside a reusable Scikit-learn pipeline so training and deployment behave consistently.",4)
    workflow_node(s,0.68,2.10,2.30,1.46,1,"Load","Read the semicolon-separated UCI CSV and inspect schema.",accent=GOLD)
    add_arrow(s,3.08,2.74)
    workflow_node(s,3.54,2.10,2.30,1.46,2,"Separate","Drop quality from X; keep quality as y.",accent=RED2)
    add_arrow(s,5.94,2.74)
    workflow_node(s,6.40,2.10,2.30,1.46,3,"Split","80% train, 20% test with random_state=42.",accent=GOLD)
    add_arrow(s,8.80,2.74)
    workflow_node(s,9.26,2.10,2.96,1.46,4,"Transform","Median imputer → StandardScaler → regressor.",accent=RED2)
    box(s,0.68,4.08,5.82,2.27,fill=PANEL,line="#39343E",radius=True)
    label(s,0.98,4.38,"WHY THIS PIPELINE",color=GOLD)
    add_bullet(s,0.98,4.79,5.00,"Imputation makes future missing inputs safe.",size=11.3)
    add_bullet(s,0.98,5.31,5.00,"Scaling standardizes feature ranges consistently.",size=11.3,accent=RED2)
    add_bullet(s,0.98,5.83,5.00,"The exact fitted steps are saved with the model.",size=11.3)
    box(s,6.82,4.08,5.85,2.27,fill=PANEL,line="#39343E",radius=True)
    label(s,7.12,4.38,"EDA CHECKS",color=RED2)
    text(s,7.12,4.83,5.05,0.53,"Distribution of labels  ·  feature correlations  ·  target relationship",size=13,color=INK,bold=True)
    text(s,7.12,5.57,4.98,0.48,"The dataset has 0 missing values. The exploratory plots are retained in artifacts/ for the report and notebook.",size=10.2,color=MUTED)

    s=prs.slides.add_slide(blank); add_bg(s); title(s,"04 · Models","Two regressors, one fair comparison","Both models see the same split and the same preprocessing. They are compared independently; they do not work together as a stacked model.",5)
    box(s,0.68,2.03,5.76,4.38,fill=PANEL,line="#39343E",radius=True)
    label(s,1.00,2.34,"MODEL 01",color=RED2)
    text(s,1.00,2.72,4.85,0.36,"Decision Tree Regressor",size=19,color=INK,bold=True,font=DISPLAY)
    text(s,1.00,3.26,4.86,0.54,"Learns a sequence of if/then splits and predicts the mean value at the final leaf.",size=11.4,color=MUTED)
    pill(s,1.00,4.06,1.34,"max_depth=5",fill=RED,color=INK,size=8.8)
    pill(s,2.48,4.06,1.54,"min_leaf=3",fill=PANEL2,color=GOLD2,size=8.8)
    text(s,1.00,4.83,4.80,0.73,"Strength: easy to explain.\nTrade-off: a single tree can be unstable.",size=11.2,color=INK)
    rule(s,1.00,5.83,1.04,RED2,height=0.04)
    text(s,1.00,6.00,4.80,0.18,"BASELINE MODEL",size=8.6,color=MUTED,bold=True)
    box(s,6.88,2.03,5.78,4.38,fill=PANEL,line="#39343E",radius=True)
    label(s,7.20,2.34,"MODEL 02 · SELECTED",color=GOLD)
    text(s,7.20,2.72,4.88,0.36,"Random Forest Regressor",size=19,color=INK,bold=True,font=DISPLAY)
    text(s,7.20,3.26,4.85,0.54,"Trains many randomized trees and averages their predictions for a more stable estimate.",size=11.4,color=MUTED)
    pill(s,7.20,4.06,1.62,"300 trees",fill=RED,color=INK,size=8.8)
    pill(s,8.98,4.06,1.54,"min_leaf=2",fill=PANEL2,color=GOLD2,size=8.8)
    text(s,7.20,4.83,4.80,0.73,"Strength: reduces variance through averaging.\nSelection rule: lowest test-set RMSE.",size=11.2,color=INK)
    rule(s,7.20,5.83,1.04,GOLD,height=0.04)
    text(s,7.20,6.00,4.80,0.18,"FINAL DEPLOYED MODEL",size=8.6,color=MUTED,bold=True)
    add_arrow(s,5.96,3.82,w=0.52,color=GOLD)

    s=prs.slides.add_slide(blank); add_bg(s); title(s,"05 · Implementation","The model is a saved, reusable pipeline","Training, evaluation, artifact export and Streamlit inference are separated so the same model is used end to end.",6)
    box(s,0.68,2.00,6.64,4.55,fill=PANEL,line="#39343E",radius=True)
    label(s,0.98,2.30,"TRAINING CORE · train_model.py",color=GOLD)
    codebox(s,0.98,2.72,6.03,3.26,[
        'def make_pipeline(model):',
        '    return Pipeline([',
        '        ("imputer", SimpleImputer(strategy="median")),',
        '        ("scaler", StandardScaler()),',
        '        ("model", model),',
        '    ])',
        '',
        'pipeline.fit(train_features, train_target)',
        'predictions = pipeline.predict(test_features)',
    ],highlight={0,1,2,3,4,7,8},size=10.1)
    text(s,0.98,6.16,5.98,0.20,"One pipeline per model keeps the comparison fair.",size=9.6,color=MUTED)
    box(s,7.62,2.00,5.05,4.55,fill=PANEL,line="#39343E",radius=True)
    label(s,7.92,2.30,"REPRODUCIBLE DELIVERY",color=RED2)
    workflow_node(s,7.92,2.77,4.45,0.78,1,"Train","python train_model.py",accent=GOLD)
    workflow_node(s,7.92,3.75,4.45,0.78,2,"Save","models/wine_quality_model.joblib",accent=RED2)
    workflow_node(s,7.92,4.73,4.45,0.78,3,"Serve","streamlit run app.py",accent=GOLD)
    text(s,7.92,5.82,4.45,0.43,"The saved file contains imputation, scaling and the selected Random Forest.\nNotebook: wine_quality_prediction.ipynb",size=9.4,color=MUTED)

    s=prs.slides.add_slide(blank); add_bg(s); title(s,"06 · Results","Random Forest wins the held-out comparison","The final model is selected by the lowest RMSE, while R² shows the proportion of target variation explained.",7)
    s.shapes.add_picture(str(ASSET/"model_comparison_dark.png"),Inches(0.68),Inches(1.93),width=Inches(6.52),height=Inches(3.30))
    box(s,7.48,1.96,5.19,3.28,fill=PANEL,line="#39343E",radius=True)
    label(s,7.78,2.25,"TEST-SET METRICS",color=GOLD)
    headers=[("MODEL",7.80,1.45),("MAE",9.28,0.68),("MSE",10.02,0.68),("RMSE",10.76,0.78),("R²",11.55,0.65)]
    for h,x,w in headers: text(s,x,2.72,w,0.18,h,size=8.1,color=MUTED,bold=True)
    rule(s,7.80,3.00,4.48,"#4A4550",height=0.018)
    rows=[("Decision Tree","0.506914","0.435915","0.660239","0.332959",RED2),("Random Forest","0.429716","0.311323","0.557963","0.523612",GOLD)]
    for i,(m,mae,mse,rmse,r2,c) in enumerate(rows):
        yy=3.26+i*0.70
        text(s,7.80,yy,1.45,0.30,m,size=9.0,color=INK,bold=(i==1))
        text(s,9.28,yy,0.68,0.30,mae,size=8.2,color=INK,bold=(i==1),align=PP_ALIGN.CENTER)
        text(s,10.02,yy,0.68,0.30,mse,size=8.2,color=INK,bold=(i==1),align=PP_ALIGN.CENTER)
        text(s,10.76,yy,0.78,0.30,rmse,size=8.2,color=INK,bold=(i==1),align=PP_ALIGN.CENTER)
        text(s,11.55,yy,0.65,0.30,r2,size=7.7,color=INK,bold=(i==1),align=PP_ALIGN.CENTER)
        if i==1: rule(s,7.80,yy+0.42,4.48,GOLD,height=0.025)
    text(s,7.80,4.82,4.35,0.24,"MAE, MSE, RMSE ↓  ·  R² ↑  ·  exact test-set values",size=8.7,color=MUTED,italic=True)
    box(s,0.68,5.58,11.99,0.75,fill=RED,line=None,radius=True)
    rich_text(s,0.96,5.78,11.35,0.31,[("DECISION  ",10,True, None,GOLD),("Random Forest is deployed because its RMSE is lower by 0.102 and its R² is higher by 0.191.",11,False,None,INK)],valign=MSO_ANCHOR.MIDDLE,margin=0)

    s=prs.slides.add_slide(blank); add_bg(s); title(s,"07 · Explainability","What the fitted forest pays attention to","Feature importance helps interpret model behavior; it describes predictive contribution, not causation.",8)
    s.shapes.add_picture(str(ASSET/"feature_importance_dark.png"),Inches(0.68),Inches(1.88),width=Inches(7.18),height=Inches(4.45))
    box(s,8.12,1.98,4.55,4.24,fill=PANEL,line="#39343E",radius=True)
    label(s,8.45,2.28,"TOP SIGNALS",color=GOLD)
    ranked=[("Alcohol","28.4%",GOLD),("Sulphates","15.9%",RED2),("Volatile acidity","11.0%",GOLD),("Total sulfur dioxide","7.8%",RED2),("Chlorides","6.7%",GOLD)]
    for i,(name,val,c) in enumerate(ranked):
        yy=2.78+i*0.54
        text(s,8.45,yy,2.40,0.24,name,size=10.6,color=INK,bold=(i<2))
        text(s,11.16,yy,0.78,0.24,val,size=10.6,color=c,bold=True,align=PP_ALIGN.RIGHT)
        rule(s,8.45,yy+0.34,3.49,"#403B45",height=0.02)
    text(s,8.45,5.70,3.55,0.30,"Alcohol is the strongest signal in this fitted model.",size=10.3,color=MUTED)

    s=prs.slides.add_slide(blank); add_bg(s); title(s,"08 · Deployment","From trained pipeline to a working Streamlit app","The app accepts the same 11 chemistry inputs, loads the saved pipeline and returns a quality estimate instantly.",9)
    box(s,0.68,2.00,5.86,4.40,fill=PANEL,line="#39343E",radius=True)
    label(s,0.98,2.30,"INFERENCE FLOW · app.py",color=GOLD)
    codebox(s,0.98,2.72,5.22,2.13,[
        'model = joblib.load(MODEL_PATH)',
        'input_frame = pd.DataFrame([input_values],',
        '    columns=metadata["features"]',
        ')',
        'prediction = model.predict(input_frame)[0]',
    ],highlight={0,1,2,4},size=10.1)
    workflow_node(s,0.98,5.11,1.48,0.88,1,"Inputs","11 fields",accent=GOLD)
    add_arrow(s,2.60,5.47,w=0.34)
    workflow_node(s,3.02,5.11,1.48,0.88,2,"Model","RF pipeline",accent=RED2)
    add_arrow(s,4.64,5.47,w=0.34)
    workflow_node(s,5.06,5.11,1.22,0.88,3,"Output","score",accent=GOLD)
    box(s,6.90,2.00,5.77,4.40,fill=BLACK,line="#3B3640",radius=True)
    label(s,7.24,2.30,"USER EXPERIENCE",color=RED2)
    text(s,7.24,2.78,4.90,0.40,"Predicted wine quality",size=13,color=MUTED)
    text(s,7.24,3.20,4.65,0.73,"5.69 / 10",size=34,color=GOLD,bold=True,font=DISPLAY)
    text(s,7.24,4.08,4.72,0.35,"Editable inputs  ·  safe training ranges  ·  instant prediction",size=10.6,color=INK)
    rule(s,7.24,4.74,1.04,RED2,height=0.045)
    text(s,7.24,5.02,4.60,0.44,"Live app and source repository are linked on the final slide.",size=11.1,color=MUTED)

    s=prs.slides.add_slide(blank); add_bg(s,ASSET/"cover_bg.png"); title(s,"09 · Conclusions","A complete, explainable ML workflow","The project demonstrates the full journey from a public dataset to a working deployed model.",10)
    box(s,0.68,2.00,7.22,3.98,fill=BLACK,line="#3B3640",radius=True,transparency=7)
    label(s,1.00,2.30,"WHAT WE BUILT",color=GOLD)
    add_bullet(s,1.00,2.80,6.25,"A reproducible regression pipeline over 1,599 red wines.",size=12)
    add_bullet(s,1.00,3.39,6.25,"A fair Decision Tree vs Random Forest comparison.",size=12,accent=RED2)
    add_bullet(s,1.00,3.98,6.25,"A selected Random Forest with test RMSE 0.558 and R² 0.524.",size=12)
    add_bullet(s,1.00,4.57,6.25,"A Streamlit interface backed by the exact saved pipeline.",size=12,accent=RED2)
    text(s,1.00,5.35,6.05,0.34,"Next: cross-validation, broader wine styles and richer sensory/context data.",size=10.5,color=MUTED,italic=True)
    box(s,8.25,2.00,4.42,3.98,fill=PANEL,line="#39343E",radius=True)
    label(s,8.57,2.30,"ACCESS",color=RED2)
    text(s,8.57,2.76,3.77,0.22,"SOURCE REPOSITORY",size=9,color=MUTED,bold=True)
    add_link(s,8.57,3.10,3.55,0.45,"github.com/SoumitraDeshpande11/\nwine-quality-prediction","https://github.com/SoumitraDeshpande11/wine-quality-prediction",color=GOLD,size=10.1)
    text(s,8.57,4.00,3.77,0.22,"LIVE STREAMLIT APP",size=9,color=MUTED,bold=True)
    add_link(s,8.57,4.33,3.58,0.66,"soumitradeshpande11-wine-quality-\nprediction-app-pnq8pq.streamlit.app","https://soumitradeshpande11-wine-quality-prediction-app-pnq8pq.streamlit.app/",color=GOLD,size=9.8)
    pill(s,8.57,5.35,1.22,"THANK YOU",fill=RED,color=INK,size=9)
    text(s,10.02,5.35,2.1,0.20,"TEAM BDA",size=10,color=GOLD,bold=True)

    prs.core_properties.title = "Wine Quality Prediction | Team BDA"
    prs.core_properties.subject = "Machine Learning Fundamentals mini project presentation"
    prs.core_properties.author = "Team BDA"
    prs.core_properties.keywords = "wine quality, machine learning, regression, random forest, streamlit"
    prs.save(OUT)
    print(OUT)

if __name__ == "__main__":
    build()
